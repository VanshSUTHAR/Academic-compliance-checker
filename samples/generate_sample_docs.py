"""Generate realistic conforming and non-conforming PDF and DOCX sample documents for testing."""

import os
from pathlib import Path
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import HRFlowable, PageBreak, Paragraph, SimpleDocTemplate, Spacer

OUT_DIR = Path(__file__).resolve().parent

GOOD_TEXT = """
Student Name: Jane Doe
Student ID: STU-2024-8891
Course Code: CS-401
Date: October 14, 2024

THE IMPACT OF ARTIFICIAL INTELLIGENCE ON AUTOMATED COMPLIANCE VERIFICATION

Abstract
This research paper investigates the efficacy of automated compliance verification frameworks within academic and professional evaluation workflows. As institutional submission rubrics grow increasingly complex, manual evaluation suffers from elevated error rates and grading inconsistencies. By deploying structural inspection algorithms alongside deterministic font geometry measurements, modern compliance engines achieve high reliability. This study reviews automated parsing architectures across 500 student submissions, demonstrating a 94 percent reduction in technical formatting penalties. The findings emphasize that automated pre-flight compliance checks empower researchers to eliminate preventable administrative marks deductions before final submission.

Introduction
Academic excellence requires adherence to scholarly conventions, formatting standards, and citation integrity. In contemporary higher education, students frequently lose substantial marks due to administrative non-compliance rather than substantive intellectual shortcomings (Smith, 2020). Previous investigations indicate that formatting guidelines such as 1-inch margins, consistent font hierarchies, and standardized citation styles cause significant student anxiety (Johnson & Lee, 2021). This paper proposes a systematic approach to verifying submission integrity through automated auditing.

Methodology
To evaluate document compliance, we established a modular evaluation pipeline comprising document layout parsing, font geometry profiling, and citation graph validation. We examined 500 academic papers formatted under APA 7th Edition standards. Document geometry was extracted via bounding box analysis, while in-text citation instances were matched against terminal bibliographies using string distance metrics (Brown et al., 2022). Quantitative compliance scores were computed across five categories: typography, structural presence, citation fidelity, metadata presence, and lexical clarity.

Results
The experimental results demonstrate that automated inspection accurately identified 98.4 percent of formatting deviations and 100 percent of missing mandatory sections. Furthermore, citation cross-referencing successfully flagged orphaned citations and unreferenced bibliography entries (Davis, 2023). Overall, the intervention improved compliance adherence from 71.2 percent to 96.8 percent across consecutive assignment submissions.

Discussion
The analysis reveals that the primary drivers of compliance failure are inconsistent margin settings and mismatched in-text citations. When automated tools provide targeted feedback before final grading, students correct technical errors proactively. These results align with broader educational literature advocating for iterative submission workflows and formative diagnostic feedback mechanisms.

Conclusion
In conclusion, automated compliance verification tools significantly reduce preventable formatting penalties in student submissions. By offering deterministic layout measurement and citation cross-checking, the framework fosters greater academic rigor and preserves grading focus on intellectual quality. Future research will explore integration with modern Learning Management Systems.

References
Brown, T., Williams, C., & Patel, K. (2022). Natural language processing for academic integrity. Journal of Educational Technology, 45(2), 112-128.
Davis, M. (2023). Document geometry and automated layout verification. Digital Publishing Quarterly, 18(4), 205-221.
Johnson, R., & Lee, S. (2021). The cognitive load of submission guidelines in higher education. Higher Education Studies, 33(1), 45-59.
Smith, A. (2020). Grading disparities and formatting compliance in university coursework. Academic Review, 29(3), 88-102.
"""

BAD_TEXT = """
Student Name: John Smith
Date: October 2024

A Quick Study on Compliance Stuff

Introduction
Everyone knows that formatting papers takes a lot of time. Students do a lot of things to try to get good grades, but sometimes teachers take points off for silly things. It is very good to check your paper before giving it to the teacher because obviously bad mistakes happen basically all the time. In this paper we did some things to see what happens when students submit papers (Miller, 2019).

Conclusion
In conclusion, students should check their papers. This stuff is very obvious and everyone knows it.

Methodology
We looked at some papers and saw that people made mistakes. Some people used the wrong fonts and wrong margins. It was basically a simple survey.

Results
The results were huge. Many papers had errors.

References
Taylor, P. (2018). Some thoughts on writing papers. General Press.
"""

def generate_good_pdf(path: Path) -> None:
    """Generate conforming APA PDF."""
    doc = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        leftMargin=1.0 * inch,
        rightMargin=1.0 * inch,
        topMargin=1.0 * inch,
        bottomMargin=1.0 * inch,
    )
    styles = getSampleStyleSheet()
    body_style = ParagraphStyle(
        "GoodBody",
        parent=styles["Normal"],
        fontName="Times-Roman",
        fontSize=12,
        leading=24,  # Double spacing
        textColor=colors.black,
    )
    title_style = ParagraphStyle(
        "GoodTitle",
        parent=styles["Heading1"],
        fontName="Times-Bold",
        fontSize=14,
        leading=28,
        alignment=1,  # Centered
    )
    heading_style = ParagraphStyle(
        "GoodHeading",
        parent=styles["Heading2"],
        fontName="Times-Bold",
        fontSize=12,
        leading=24,
        alignment=1,  # Centered for APA section titles
    )

    story = []
    lines = [ln.strip() for ln in GOOD_TEXT.split("\n") if ln.strip()]

    # Metadata lines
    for line in lines[:4]:
        story.append(Paragraph(line, body_style))
    story.append(Spacer(1, 14))

    # Title
    story.append(Paragraph(lines[4], title_style))
    story.append(Spacer(1, 14))

    # Content
    i = 5
    while i < len(lines):
        line = lines[i]
        if line in ["Abstract", "Introduction", "Methodology", "Results", "Discussion", "Conclusion", "References"]:
            story.append(Spacer(1, 10))
            story.append(Paragraph(line, heading_style))
            story.append(Spacer(1, 6))
        else:
            story.append(Paragraph(line, body_style))
            story.append(Spacer(1, 6))
        i += 1

    doc.build(story)

def generate_bad_pdf(path: Path) -> None:
    """Generate non-conforming PDF (wrong margins 0.5", Helvetica 10pt, missing sections, missing metadata)."""
    doc = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        leftMargin=0.5 * inch,
        rightMargin=0.5 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
    )
    styles = getSampleStyleSheet()
    body_style = ParagraphStyle(
        "BadBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=13,  # Single spacing
        textColor=colors.black,
    )
    story = []
    for line in BAD_TEXT.split("\n"):
        if line.strip():
            story.append(Paragraph(line.strip(), body_style))
            story.append(Spacer(1, 4))
    doc.build(story)

def generate_good_docx(path: Path) -> None:
    """Generate conforming APA DOCX."""
    doc = docx.Document()
    # Set 1-inch margins
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    # Set default style font
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    lines = [ln.strip() for ln in GOOD_TEXT.split("\n") if ln.strip()]
    for line in lines[:4]:
        p = doc.add_paragraph(line)
        p.paragraph_format.line_spacing = 2.0

    p_title = doc.add_paragraph(lines[4])
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.line_spacing = 2.0
    for run in p_title.runs:
        run.bold = True

    i = 5
    while i < len(lines):
        line = lines[i]
        if line in ["Abstract", "Introduction", "Methodology", "Results", "Discussion", "Conclusion", "References"]:
            h = doc.add_heading(line, level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            h.paragraph_format.line_spacing = 2.0
            for r in h.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)
                r.font.color.rgb = RGBColor(0, 0, 0)
        else:
            p = doc.add_paragraph(line)
            p.paragraph_format.line_spacing = 2.0
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)
        i += 1

    doc.save(str(path))

def generate_bad_docx(path: Path) -> None:
    """Generate non-conforming DOCX (0.5 inch margins, Arial 10pt, missing metadata)."""
    doc = docx.Document()
    for sec in doc.sections:
        sec.top_margin = Inches(0.5)
        sec.bottom_margin = Inches(0.5)
        sec.left_margin = Inches(0.5)
        sec.right_margin = Inches(0.5)

    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    for line in BAD_TEXT.split("\n"):
        if line.strip():
            p = doc.add_paragraph(line.strip())
            p.paragraph_format.line_spacing = 1.0
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(10)

    doc.save(str(path))

def generate_all_samples() -> None:
    """Generate all test documents in samples directory."""
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    print("Generating sample documents...")
    generate_good_pdf(OUT_DIR / "sample_assignment_good.pdf")
    generate_bad_pdf(OUT_DIR / "sample_assignment_bad.pdf")
    generate_good_docx(OUT_DIR / "sample_assignment_good.docx")
    generate_bad_docx(OUT_DIR / "sample_assignment_bad.docx")
    print(f"Sample documents successfully generated in {OUT_DIR}")

if __name__ == "__main__":
    generate_all_samples()
